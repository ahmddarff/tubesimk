import os, io, openpyxl
from datetime import datetime, timedelta
from flask import Blueprint, render_template, request, jsonify, current_app, Response
from flask_login import login_required, current_user 
from models import User, CafeSetting, OperationalHour, Menu, Category, Table, Order, OrderItem
from sqlalchemy import func
from extensions import db
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename
from openpyxl.styles import Font, PatternFill, Alignment
from utils import *

owner_bp = Blueprint('owner', __name__)

# ── 1. DASHBOARD ENDPOINT ─────────────────────────────────────────────────────
@owner_bp.route('/dashboard')
@login_required
@role_required('owner')
def dashboard():
    hari_ini = datetime.utcnow().date()
    kemarin = hari_ini - timedelta(days=1)

    penjualan_hari_ini = db.session.query(func.sum(Order.total_amount)).filter(Order.payment_status == 'paid', func.date(Order.created_at) == hari_ini).scalar() or 0
    penjualan_kemarin = db.session.query(func.sum(Order.total_amount)).filter(Order.payment_status == 'paid', func.date(Order.created_at) == kemarin).scalar() or 0
    tren_penjualan = ((penjualan_hari_ini - penjualan_kemarin) / penjualan_kemarin * 100) if penjualan_kemarin > 0 else (100 if penjualan_hari_ini > 0 else 0)

    order_hari_ini = Order.query.filter(Order.payment_status == 'paid', func.date(Order.created_at) == hari_ini).count()
    order_kemarin = Order.query.filter(Order.payment_status == 'paid', func.date(Order.created_at) == kemarin).count()
    tren_order = ((order_hari_ini - order_kemarin) / order_kemarin * 100) if order_kemarin > 0 else (100 if order_hari_ini > 0 else 0)

    menu_terlaris_query = db.session.query(Menu.name, func.sum(OrderItem.qty).label('total_qty'))\
        .join(OrderItem, OrderItem.menu_id == Menu.id).join(Order, OrderItem.order_id == Order.id)\
        .filter(Order.payment_status == 'paid', func.date(Order.created_at) == hari_ini)\
        .group_by(Menu.name).order_by(func.sum(OrderItem.qty).desc()).first()
        
    menu_terlaris_name = menu_terlaris_query[0] if menu_terlaris_query else "Belum Ada"
    menu_terlaris_qty = menu_terlaris_query[1] if menu_terlaris_query else 0

    tujuh_hari_lalu = hari_ini - timedelta(days=6)

    category_data = db.session.query(Category.name, func.sum(OrderItem.qty))\
        .join(Menu, Menu.category_id == Category.id).join(OrderItem, OrderItem.menu_id == Menu.id).join(Order, OrderItem.order_id == Order.id)\
        .filter(Order.payment_status == 'paid', Order.created_at >= tujuh_hari_lalu)\
        .group_by(Category.name).all()
    chart_categories = [item[0] for item in category_data]
    chart_category_qtys = [int(item[1]) for item in category_data]

    pendapatan_harian = db.session.query(
        func.date(Order.created_at).label('tanggal'),
        func.sum(Order.total_amount).label('total')
    ).filter(Order.payment_status == 'paid', Order.created_at >= tujuh_hari_lalu)\
     .group_by(func.date(Order.created_at)).order_by(func.date(Order.created_at)).all()
     
    line_labels = [p.tanggal.strftime('%a') for p in pendapatan_harian] 
    line_values = [int(p.total) for p in pendapatan_harian]

    if not line_labels:
        line_labels = ['Sen', 'Sel', 'Rab', 'Kam', 'Jum', 'Sab', 'Min']
        line_values = [0, 0, 0, 0, 0, 0, 0]

    return render_template(
        'owner/dashboard.html',
        total_penjualan=f"Rp {penjualan_hari_ini:,}".replace(",", "."),
        tren_penjualan=round(abs(tren_penjualan), 1),
        is_penjualan_naik=(tren_penjualan >= 0),
        total_order=order_hari_ini,
        tren_order=round(abs(tren_order), 1),
        is_order_naik=(tren_order >= 0),
        menu_terlaris=menu_terlaris_name,
        menu_terlaris_qty=menu_terlaris_qty,
        line_labels=line_labels,
        line_values=line_values,
        chart_categories=chart_categories,
        chart_category_qtys=chart_category_qtys
    )

@owner_bp.route('/api/dashboard/filter', methods=['POST'])
@login_required
@role_required('owner')
def api_dashboard_filter():
    data = request.json
    rentang = data.get('rentang')
    
    # Hitung waktu (Logika WIB-UTC yang sudah kita buat)
    sekarang_wib = datetime.utcnow() + timedelta(hours=7)
    hari_ini_wib = sekarang_wib.date()
    
    if rentang == 'Bulan Ini':
        start_date_wib = hari_ini_wib.replace(day=1)
        prev_end_wib = start_date_wib - timedelta(days=1)
        prev_start_wib = prev_end_wib.replace(day=1)
    elif rentang == '7 Hari Terakhir':
        start_date_wib = hari_ini_wib - timedelta(days=6)
        prev_start_wib = start_date_wib - timedelta(days=7)
        prev_end_wib = start_date_wib - timedelta(days=1)
    else:
        start_date_wib = hari_ini_wib
        prev_start_wib = hari_ini_wib - timedelta(days=1)
        prev_end_wib = prev_start_wib

    start_utc = datetime.combine(start_date_wib, datetime.min.time()) - timedelta(hours=7)
    end_utc = datetime.combine(hari_ini_wib, datetime.max.time()) - timedelta(hours=7)
    prev_start_utc = datetime.combine(prev_start_wib, datetime.min.time()) - timedelta(hours=7)
    prev_end_utc = datetime.combine(prev_end_wib, datetime.max.time()) - timedelta(hours=7)

    # 1. Penjualan
    curr_penjualan = db.session.query(func.sum(Order.total_amount)).filter(Order.payment_status == 'paid', Order.created_at >= start_utc, Order.created_at <= end_utc).scalar() or 0
    prev_penjualan = db.session.query(func.sum(Order.total_amount)).filter(Order.payment_status == 'paid', Order.created_at >= prev_start_utc, Order.created_at <= prev_end_utc).scalar() or 0
    tren_penjualan = ((curr_penjualan - prev_penjualan) / prev_penjualan * 100) if prev_penjualan > 0 else (100 if curr_penjualan > 0 else 0)

    # 2. Order
    curr_order = Order.query.filter(Order.payment_status == 'paid', Order.created_at >= start_utc, Order.created_at <= end_utc).count()
    prev_order = Order.query.filter(Order.payment_status == 'paid', Order.created_at >= prev_start_utc, Order.created_at <= prev_end_utc).count()
    tren_order = ((curr_order - prev_order) / prev_order * 100) if prev_order > 0 else (100 if curr_order > 0 else 0)

    # 3. Menu Terlaris
    menu_q = db.session.query(Menu.name, func.sum(OrderItem.qty).label('qty'))\
        .join(OrderItem, OrderItem.menu_id == Menu.id).join(Order, OrderItem.order_id == Order.id)\
        .filter(Order.payment_status == 'paid', Order.created_at >= start_utc, Order.created_at <= end_utc)\
        .group_by(Menu.name).order_by(func.sum(OrderItem.qty).desc()).first()

    # 4. Kategori
    cat_data = db.session.query(Category.name, func.sum(OrderItem.qty))\
        .join(Menu, Menu.category_id == Category.id).join(OrderItem, OrderItem.menu_id == Menu.id).join(Order, OrderItem.order_id == Order.id)\
        .filter(Order.payment_status == 'paid', Order.created_at >= start_utc, Order.created_at <= end_utc)\
        .group_by(Category.name).all()

    return jsonify({
        "sukses": True,
        "penjualan": {"total": f"Rp {curr_penjualan:,}".replace(",", "."), "tren": round(abs(tren_penjualan), 1), "naik": tren_penjualan >= 0},
        "order": {"total": curr_order, "tren": round(abs(tren_order), 1), "naik": tren_order >= 0},
        "menu": {"nama": menu_q[0] if menu_q else "Belum Ada", "qty": int(menu_q[1]) if menu_q else 0},
        "kategori": {"labels": [i[0] for i in cat_data], "values": [int(i[1]) for i in cat_data]}
    })

# ── 2. CORE MANAGEMENT PAGES (DINAMIS DB) ────────────────────────────────────
@owner_bp.route('/manajemen-menu')
@login_required
@role_required('owner')
def manajemen_menu():
    menus = Menu.query.all()
    categories = Category.query.all()
    
    for m in menus:
        all_ratings = [oi.review.rating for oi in m.order_items if oi.review]
        if all_ratings:
            m.avg_rating = round(sum(all_ratings) / len(all_ratings), 1)
            m.total_reviews = len(all_ratings)
        else:
            m.avg_rating = 0.0
            m.total_reviews = 0

    return render_template('owner/manajemen-menu.html', menu_list=menus, categories=categories)

@owner_bp.route('/manajemen-meja')
@login_required
@role_required('owner')
def manajemen_meja():
    tables = Table.query.order_by(Table.table_number).all()
    return render_template('owner/manajemen-meja.html', tables=tables)

@owner_bp.route('/manajemen-staff')
@login_required
@role_required('owner')
def manajemen_kasir():
    # 1. Ambil Data Kasir
    kasir_users = User.query.filter_by(role='kasir').all()
    kasir_list = []
    for k in kasir_users:
        total_tx = Order.query.filter_by(cashier_id=k.id, payment_status='paid').count()
        total_sales = db.session.query(func.sum(Order.total_amount)).filter_by(cashier_id=k.id, payment_status='paid').scalar() or 0
        
        kasir_list.append({
            "id": k.id,
            "nama": k.name,
            "username": k.username,
            "status": "online" if k.is_active else "offline",
            "total_transaksi": total_tx,
            "total_penjualan": total_sales
        })
        
    # 2. Ambil Data Koki (SEKARANG DINAMIS)
    koki_users = User.query.filter_by(role='koki').all()
    koki_list = []
    for c in koki_users:
        # Tarik semua menu yang dikerjakan koki ini dan berstatus selesai (ready/served)
        finished_items = OrderItem.query.filter_by(koki_id=c.id).filter(
            OrderItem.item_status.in_(['ready', 'served'])
        ).all()
        
        total_pesanan = len(finished_items)
        total_durasi_menit = 0
        valid_items_count = 0
        
        for item in finished_items:
            # Pastikan kedua timestamp (mulai dan selesai) ada untuk menghindari eror
            if item.prepared_at and item.ready_at:
                durasi_detik = (item.ready_at - item.prepared_at).total_seconds()
                total_durasi_menit += (durasi_detik / 60.0)
                valid_items_count += 1
                
        # Hitung rata-rata jika ada pesanan yang valid
        if valid_items_count > 0:
            avg_speed_menit = round(total_durasi_menit / valid_items_count)
            avg_speed_str = f"{avg_speed_menit} Menit"
        else:
            avg_speed_str = "-"

        koki_list.append({
            "id": c.id,
            "nama": c.name,
            "username": c.username,
            "total_pesanan": total_pesanan,  
            "avg_speed": avg_speed_str,    
            "status": "online" if c.is_active else "offline"
        })

    # 3. Hitung Statistik Ringkasan
    kasir_online = sum(1 for k in kasir_list if k['status'] == 'online')
    koki_online = sum(1 for c in koki_list if c['status'] == 'online')
    total_staff = len(kasir_list) + len(koki_list)

    return render_template('owner/manajemen-staff.html',
        kasir_list=kasir_list,
        koki_list=koki_list,
        kasir_online=kasir_online,
        koki_online=koki_online,
        total_staff=total_staff
    )

@owner_bp.route('/laporan-penjualan')
@login_required
@role_required('owner')
def laporan_penjualan():
    # 1. Hitung Statistik Utama (Hanya pesanan LUNAS)
    paid_orders = Order.query.filter_by(payment_status='paid').all()
    
    total_penjualan = sum(o.total_amount for o in paid_orders)
    total_order = len(paid_orders)
    rata_rata = total_penjualan / total_order if total_order > 0 else 0

    # 2. Data Grafik Penjualan (7 Hari Terakhir)
    hari_ini = datetime.utcnow().date()
    tujuh_hari_lalu = hari_ini - timedelta(days=6)
    
    pendapatan_harian = db.session.query(
        func.date(Order.created_at).label('tanggal'),
        func.sum(Order.total_amount).label('total')
    ).filter(Order.payment_status == 'paid', Order.created_at >= tujuh_hari_lalu)\
        .group_by(func.date(Order.created_at))\
        .order_by(func.date(Order.created_at)).all()
    
    chart_labels = [p.tanggal.strftime('%d %b') for p in pendapatan_harian]
    chart_data = [int(p.total) for p in pendapatan_harian]

    if not chart_labels:
        chart_labels = ['Sen', 'Sel', 'Rab', 'Kam', 'Jum', 'Sab', 'Min']
        chart_data = [0, 0, 0, 0, 0, 0, 0]

    # 3. Kueri Database Utama
    orders_db = Order.query.filter(
        db.or_(
            db.and_(Order.payment_status == 'paid', Order.order_status == 'served', Order.table_id.is_(None)),
            Order.payment_status == 'cancelled'
        )
    ).order_by(Order.created_at.desc()).all()
    
    data_transaksi = []
    
    for t in orders_db:
        items_list = []
        for item in t.items:
            items_list.append({
                'nama': item.menu.name if item.menu else 'Item Terhapus',
                'harga': item.price_at_order,
                'jumlah': item.qty,
                'subtotal': item.price_at_order * item.qty,
                'note': item.notes or ''
            })
            
        data_transaksi.append({
            'order_number': t.order_number,
            'tanggal': f"{format_tanggal_lokal(t.created_at)}, {format_waktu_lokal(t.created_at)}",
            'tanggal_mentah': format_tanggal_mentah(t.created_at),
            'sumber': 'App Mandiri' if t.user_id else 'Kasir',
            'nama_kasir': t.cashier.name if t.cashier else 'Self-Service',
            'metode': t.payment_method.upper() if t.payment_method else '-',
            'total': t.total_amount,
            'status': t.payment_status,
            'alasan_batal': t.cancellation_reason or '',
            'items': items_list
        })
    
    return render_template('owner/laporan-penjualan.html', 
                            transaksi_list=data_transaksi, # Mengirim data yang sudah rapi
                            total_penjualan=total_penjualan,
                            total_order=total_order,
                            rata_rata=rata_rata,
                            chart_labels=chart_labels,
                            chart_data=chart_data)

@owner_bp.route('/pengaturan')
@login_required
@role_required('owner')
def pengaturan():
    cafe_info = CafeSetting.query.first()
    jam_db = OperationalHour.query.all()
    urutan_hari = {"Senin": 1, "Selasa": 2, "Rabu": 3, "Kamis": 4, "Jumat": 5, "Sabtu": 6, "Minggu": 7}
    jam_db.sort(key=lambda x: urutan_hari.get(x.day_of_week, 8))

    return render_template('owner/pengaturan.html', cafe_info=cafe_info, jam_operasional=jam_db)

# ── 3. KATEGORI & MENU APIS ───────────────────────────────────────────────────
@owner_bp.route('/api/tambah-kategori', methods=['POST'])
@login_required
@role_required('owner')
def tambah_kategori():
    nama = request.form.get('nama')
    tipe = request.form.get('type', 'food')
    if not nama:
        return jsonify({"success": False, "message": "Nama kategori tidak boleh kosong."})
    
    if Category.query.filter_by(name=nama).first():
        return jsonify({"success": False, "message": "Kategori tersebut sudah ada."})
    
    try:
        db.session.add(Category(name=nama, type=tipe))
        db.session.commit()
        return jsonify({"success": True, "message": "Kategori baru berhasil ditambahkan!"})
    except Exception as e:
        db.session.rollback()
        return jsonify({"success": False, "message": str(e)})

@owner_bp.route('/api/edit-kategori/<int:id>', methods=['POST'])
@login_required
@role_required('owner')
def edit_kategori(id):
    cat = db.session.get(Category, id)
    if not cat:
        return jsonify({"success": False, "message": "Kategori tidak ditemukan."})
    
    nama_baru = request.form.get('nama')
    tipe_baru = request.form.get('type')

    if not nama_baru:
        return jsonify({"success": False, "message": "Nama kategori tidak boleh kosong."})
        
    if Category.query.filter(Category.name == nama_baru, Category.id != id).first():
        return jsonify({"success": False, "message": "Nama kategori ini sudah digunakan."})
        
    try:
        cat.name = nama_baru
        if tipe_baru:
            cat.type = tipe_baru
        db.session.commit()
        return jsonify({"success": True, "message": "Nama kategori berhasil diubah!"})
    except Exception as e:
        db.session.rollback()
        return jsonify({"success": False, "message": str(e)})

@owner_bp.route('/api/hapus-kategori/<int:id>', methods=['POST'])
@login_required
@role_required('owner')
def hapus_kategori(id):
    cat = db.session.get(Category, id)
    if not cat:
        return jsonify({"success": False, "message": "Kategori tidak ditemukan."})
        
    if cat.menus:
        return jsonify({"success": False, "message": f"Gagal! Kategori digunakan oleh {len(cat.menus)} menu."})
        
    try:
        db.session.delete(cat)
        db.session.commit()
        return jsonify({"success": True, "message": "Kategori berhasil dihapus!"})
    except Exception as e:
        db.session.rollback()
        return jsonify({"success": False, "message": str(e)})

@owner_bp.route('/api/tambah-menu', methods=['POST'])
@login_required
@role_required('owner')
def tambah_menu():
    try:
        nama = request.form.get('nama')
        kategori_nama = request.form.get('kategori')
        harga = request.form.get('harga')
        stok_raw = request.form.get('stok')
        deskripsi = request.form.get('deskripsi')

        stok = int(stok_raw) if (stok_raw and stok_raw.strip() != "") else None
        category = Category.query.filter_by(name=kategori_nama).first()
        if not category:
            return jsonify({"success": False, "message": "Kategori tidak ditemukan."})

        foto = request.files.get('foto')
        img_path = None

        if foto and foto.filename != '':
            filename = secure_filename(f"{datetime.now().strftime('%Y%m%d%H%M%S')}_{foto.filename}")
            upload_path = os.path.join(current_app.root_path, 'static/uploads/menu')
            if not os.path.exists(upload_path):
                os.makedirs(upload_path)
            foto.save(os.path.join(upload_path, filename))
            img_path = f"uploads/menu/{filename}"

        new_menu = Menu(
            name=nama, category_id=category.id, price=int(harga or 0),
            stock=stok, description=deskripsi, image_url=img_path, is_available=True
        )
        db.session.add(new_menu)
        db.session.commit()
        return jsonify({"success": True, "message": "Menu berhasil ditambahkan!"})
    except Exception as e:
        db.session.rollback()
        return jsonify({"success": False, "message": str(e)})

@owner_bp.route('/api/toggle-menu-status/<int:menu_id>', methods=['POST'])
@login_required
@role_required('owner')
def toggle_menu_status(menu_id):
    data = request.json
    menu = db.session.get(Menu, menu_id)
    if menu:
        menu.is_available = data.get("status")
        try:
            db.session.commit()
            return jsonify({"success": True, "message": "Status menu diperbarui!"})
        except:
            db.session.rollback()
            return jsonify({"success": False, "message": "Gagal memperbarui status."})
    return jsonify({"success": False, "message": "Menu tidak ditemukan."})

@owner_bp.route('/api/edit-menu/<int:menu_id>', methods=['POST'])
@login_required
@role_required('owner')
def edit_menu(menu_id):
    try:
        menu = db.session.get(Menu, menu_id)
        if not menu:
            return jsonify({"success": False, "message": "Menu tidak ditemukan."})

        menu.name = request.form.get('nama')
        menu.price = int(request.form.get('harga') or 0)
        menu.description = request.form.get('deskripsi')
        
        status_str = request.form.get('status')
        menu.is_available = True if status_str == 'true' else False
        
        stok_raw = request.form.get('stok')
        menu.stock = int(stok_raw) if (stok_raw and stok_raw.strip() != "") else None

        kategori_nama = request.form.get('kategori')
        category = Category.query.filter_by(name=kategori_nama).first()
        if category:
            menu.category_id = category.id

        if 'foto' in request.files:
            foto = request.files['foto']
            if foto and foto.filename != '':
                filename = secure_filename(f"{datetime.now().strftime('%Y%m%d%H%M%S')}_{foto.filename}")
                upload_path = os.path.join(current_app.root_path, 'static/uploads/menu')
                if not os.path.exists(upload_path):
                    os.makedirs(upload_path)
                foto.save(os.path.join(upload_path, filename))
                menu.image_url = f"uploads/menu/{filename}"

        db.session.commit()
        return jsonify({"success": True, "message": "Perubahan berhasil disimpan!"})
    except Exception as e:
        db.session.rollback()
        return jsonify({"success": False, "message": str(e)})

@owner_bp.route('/api/update-stok/<int:menu_id>', methods=['POST'])
@login_required
@role_required('owner')
def update_stok_cepat(menu_id):
    data = request.json
    try:
        menu = db.session.get(Menu, menu_id)
        if not menu:
            return jsonify({"success": False, "message": "Menu tidak ditemukan."})
            
        menu.stock = int(data.get('stok'))
        
        if menu.stock == 0:
            menu.is_available = False
        elif menu.stock > 0:
            menu.is_available = True

        db.session.commit()
        return jsonify({"success": True, "message": "Stok diperbarui!"})
    except Exception as e:
        db.session.rollback()
        return jsonify({"success": False, "message": "Gagal update stok."})

@owner_bp.route('/api/menu-reviews/<int:menu_id>', methods=['GET'])
@login_required
@role_required('owner')
def get_menu_reviews(menu_id):
    menu = db.session.get(Menu, menu_id)
    if not menu:
        return jsonify({"success": False, "message": "Menu tidak ditemukan"})

    reviews_data = []
    for oi in menu.order_items:
        if oi.review:
            customer_name = oi.order.customer_name or (oi.order.customer.name if oi.order.customer else "Tamu Anonim")
            reviews_data.append({
                "rating": oi.review.rating,
                "comment": oi.review.comment or "Tidak ada komentar tulisan.",
                "date": oi.review.created_at.strftime("%d %b %Y"),
                "customer": customer_name
            })

    reviews_data.sort(key=lambda x: datetime.strptime(x['date'], "%d %b %Y"), reverse=True)
    return jsonify({"success": True, "reviews": reviews_data})

# ── 4. MEJA APIS ──────────────────────────────────────────────────────────────
@owner_bp.route('/api/tambah-meja', methods=['POST'])
@login_required
@role_required('owner')
def tambah_meja():
    nomor = request.form.get('nomor')
    kapasitas = request.form.get('kapasitas')

    if not nomor or not kapasitas:
        return jsonify({"success": False, "message": "Nomor dan kapasitas wajib diisi!"})
    
    if Table.query.filter_by(table_number=nomor).first():
        return jsonify({"success": False, "message": "Nomor meja ini sudah terdaftar."})

    try:
        db.session.add(Table(table_number=nomor, capacity=int(kapasitas), is_available=True))
        db.session.commit()
        return jsonify({"success": True, "message": "Meja berhasil ditambahkan!"})
    except Exception as e:
        db.session.rollback()
        return jsonify({"success": False, "message": str(e)})

@owner_bp.route('/api/edit-meja/<int:id>', methods=['POST'])
@login_required
@role_required('owner')
def edit_meja(id):
    table = db.session.get(Table, id)
    if not table:
        return jsonify({"success": False, "message": "Meja tidak ditemukan."})

    if not table.is_available:
        return jsonify({"success": False, "message": f"Gagal! Meja {table.table_number} sedang terisi/digunakan."})

    nomor_baru = request.form.get('nomor')
    kapasitas_baru = request.form.get('kapasitas')
    
    if nomor_baru != table.table_number and Table.query.filter_by(table_number=nomor_baru).first():
        return jsonify({"success": False, "message": "Nomor meja tersebut sudah digunakan."})

    try:
        table.table_number = nomor_baru
        table.capacity = int(kapasitas_baru)
        db.session.commit()
        return jsonify({"success": True, "message": "Data meja berhasil diperbarui!"})
    except Exception as e:
        db.session.rollback()
        return jsonify({"success": False, "message": str(e)})

@owner_bp.route('/api/hapus-meja/<int:id>', methods=['POST'])
@login_required
@role_required('owner')
def hapus_meja(id):
    table = db.session.get(Table, id)
    if not table.is_available:
        return jsonify({"success": False, "message": f"Gagal! Meja {table.table_number} sedang terisi oleh tamu."})
    
    try:
        db.session.delete(table)
        db.session.commit()
        return jsonify({"success": True, "message": f"Meja {table.table_number} berhasil dihapus!"})
    except Exception as e:
        db.session.rollback()
        return jsonify({"success": False, "message": "Gagal menghapus meja."})

# ── 5. STAFF / KASIR APIS (DINAMIS DATABASE) ──────────────────────────────────
@owner_bp.route('/api/tambah-staff', methods=['POST'])
@login_required
@role_required('owner')
def tambah_staff():
    data = request.json
    nama = data.get("nama")
    username = data.get("username") or nama.lower().replace(" ", "")
    password = data.get("password") or "123456" 
    jabatan = data.get("jabatan", "kasir").lower() # ✅ Ambil pilihan jabatan dari Frontend
    
    if User.query.filter_by(username=username).first():
        return jsonify({"success": False, "message": "Username staff sudah terdaftar!"})
        
    new_staff = User(
        name=nama, 
        username=username,
        password=generate_password_hash(password),
        role=jabatan, # ✅ Simpan sesuai pilihan dropdown
        is_active=True
    )
    try:
        db.session.add(new_staff)
        db.session.commit()
        return jsonify({"success": True, "message": f"Staff '{nama}' berhasil didaftarkan sebagai {jabatan.capitalize()}!"})
    except Exception as e:
        db.session.rollback()
        return jsonify({"success": False, "message": str(e)})

@owner_bp.route('/api/toggle-kasir-status/<int:kasir_id>', methods=['POST'])
@login_required
@role_required('owner')
def toggle_kasir_status(kasir_id):
    # ✅ SEKARANG DINAMIS: Mengubah keaktifan user kasir di database
    data = request.json
    user = db.session.get(User, kasir_id)
    if user:
        status_val = data.get("status")
        user.is_active = (status_val == 'online') if status_val in ['online', 'offline'] else bool(status_val)
        db.session.commit()
        return jsonify({"success": True, "message": "Status keaktifan staf diperbarui!"})
    return jsonify({"success": False, "message": "Staff tidak ditemukan."})

@owner_bp.route('/api/edit-staff/<int:id>', methods=['POST'])
@login_required
@role_required('owner')
def edit_staff(id):
    staff = User.query.get_or_404(id)
    data = request.json
    
    nama_baru = data.get("nama")
    jabatan_baru = data.get("jabatan", "kasir").lower()
    is_active_status = data.get("isActive", True) # Menangkap status boolean checkbox
    password_baru = data.get("password")
    
    staff.name = nama_baru
    staff.role = jabatan_baru
    staff.is_active = is_active_status # Mengupdate kolom keaktifan di database

    if password_baru and password_baru.strip() != "":
        staff.password = generate_password_hash(password_baru)
    
    try:
        db.session.commit()
        return jsonify({"success": True, "message": f"Profil {nama_baru} berhasil diperbarui!"})
    except Exception as e:
        db.session.rollback()
        return jsonify({"success": False, "message": "Gagal menyimpan perubahan ke database."})

# ── 6. Laporan APIS ───────────────────────────────────────────────────────────
@owner_bp.route('/api/export-excel', methods=['POST'])
@login_required
@role_required('owner')
def export_excel():
    data = request.json
    
    # 1. Buat Buku Kerja (Workbook) Excel Asli
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Laporan Penjualan"
    
    # 2. Definisikan Gaya (Styling) untuk Header
    header_fill = PatternFill(start_color="0022AA", end_color="0022AA", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    align_center = Alignment(horizontal="center", vertical="center")
    
    headers = [
        'Nomor Order', 'Tanggal Transaksi', 'Sumber', 'Penerima Dana', 
        'Metode Pembayaran', 'Total Transaksi (Rp)', 'Status Final', 'Catatan Batal'
    ]
    
    # 3. Tulis Header dan Terapkan Gaya
    ws.append(headers)
    for col_num, cell in enumerate(ws[1], 1):
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center

    # 4. Looping Baris Data
    for t in data:
        ws.append([
            t.get('order_number', ''),
            t.get('tanggal', ''),
            t.get('sumber', ''),
            t.get('nama_kasir', ''),
            t.get('metode', ''),
            int(t.get('total', 0)),  # Pastikan berupa angka (integer)
            str(t.get('status', '')).upper(),
            t.get('alasan_batal', '')
        ])
        
    # 5. (Opsional) Rapikan Lebar Kolom Otomatis
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter # Dapatkan huruf kolom (A, B, C...)
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = (max_length + 2)
        ws.column_dimensions[column].width = adjusted_width

    # 6. Simpan File ke dalam Memori RAM (BytesIO)
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    # 7. Susun File Download dengan ekstensi .xlsx
    timestamp = datetime.now().strftime("%Y%m%d_%H%M")
    nama_file = f"Laporan_Terralog_{timestamp}.xlsx"
    
    return Response(
        output.getvalue(),
        # Mimetype biner khusus .xlsx
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={"Content-Disposition": f"attachment;filename={nama_file}"}
    )

@owner_bp.route('/cetak-laporan')
@login_required
@role_required('owner')
def cetak_laporan():
    # Ambil info cafe untuk kop surat (opsional, pastikan model CafeSetting sudah di-import)
    cafe_info = CafeSetting.query.first()
    return render_template('owner/includes/cetak_laporan.html', cafe=cafe_info)

# ── 7. CAFE SETTINGS APIS ─────────────────────────────────────────────────────
@owner_bp.route('/api/update-profil-cafe', methods=['POST'])
@login_required
@role_required('owner')
def update_profil_cafe():
    nama = request.form.get("nama")
    telp = request.form.get("telp")
    alamat = request.form.get("alamat")
    email = request.form.get("email")
    
    cafe_info = CafeSetting.query.first()
    if cafe_info:
        cafe_info.cafe_name = nama or cafe_info.cafe_name
        cafe_info.phone = telp or cafe_info.phone
        cafe_info.address = alamat or cafe_info.address
        cafe_info.email = email or cafe_info.email
        cafe_info.reservation_buffer_time = int(request.form.get('reservation_buffer_time', 90))
        cafe_info.table_clearance_time = int(request.form.get('table_clearance_time', 15))
        
        if 'logo' in request.files:
            file = request.files['logo']
            if file and file.filename != '':
                if cafe_info.logo:
                    old_path = os.path.join(current_app.root_path, 'static/images', cafe_info.logo)
                    if os.path.exists(old_path) and os.path.isfile(old_path):
                        try: os.remove(old_path)
                        except: pass
                filename = secure_filename(file.filename)
                unique_filename = f"{datetime.now().strftime('%Y%m%d%H%M%S')}_{filename}"
                file.save(os.path.join(current_app.root_path, 'static/images', unique_filename))
                cafe_info.logo = unique_filename

        try:
            db.session.commit()
            return jsonify({"success": True, "message": "Profil cafe berhasil disimpan!"})
        except:
            db.session.rollback()
            return jsonify({"success": False, "message": "Gagal menyimpan ke database."})
    return jsonify({"success": False, "message": "Data pengaturan tidak ditemukan."})

@owner_bp.route('/api/update-akun', methods=['POST'])
@login_required
@role_required('owner')
def update_akun():
    user = db.session.get(User, current_user.id)
    if user:
        user.name = request.form.get("nama") or user.name
        user.username = request.form.get("username") or user.username
        user.email = request.form.get("email") or user.email
        user.phone = request.form.get("no_hp") or user.phone
        
        if 'photo' in request.files:
            file = request.files['photo']
            if file and file.filename != '':
                upload_path = os.path.join(current_app.root_path, 'static/uploads/profile')
                if not os.path.exists(upload_path):
                    os.makedirs(upload_path)
                if user.photo:
                    old_path = os.path.join(current_app.root_path, 'static', user.photo)
                    if os.path.exists(old_path) and os.path.isfile(old_path):
                        try: os.remove(old_path)
                        except: pass
                filename = secure_filename(file.filename)
                unique_filename = f"{user.id}_{datetime.now().strftime('%Y%m%d%H%M%S')}_{filename}"
                file.save(os.path.join(upload_path, unique_filename))
                user.photo = f"uploads/profile/{unique_filename}"

        try:
            db.session.commit()
            return jsonify({"success": True, "message": "Akun berhasil diperbarui!"})
        except:
            db.session.rollback()
            return jsonify({"success": False, "message": "Gagal memperbarui database."})
    return jsonify({"success": False, "message": "Pengguna tidak ditemukan."})

@owner_bp.route('/api/update-password', methods=['POST'])
@login_required
@role_required('owner')
def update_password():
    data = request.json
    user = db.session.get(User, current_user.id)
    if not check_password_hash(user.password, data.get("password_lama")):
        return jsonify({"success": False, "message": "Kata sandi lama salah!"})
    
    user.password = generate_password_hash(data.get("password_baru"))
    try:
        db.session.commit()
        return jsonify({"success": True, "message": "Kata sandi berhasil diperbarui!"})
    except:
        db.session.rollback()
        return jsonify({"success": False, "message": "Gagal menyimpan kata sandi."})
    
@owner_bp.route('/api/toggle-cafe-status', methods=['POST'])
@login_required
@role_required('owner')
def toggle_cafe_status():
    data = request.json
    cafe = CafeSetting.query.first()
    if cafe:
        cafe.is_open = data.get("status")
        try:
            db.session.commit()
            return jsonify({"success": True, "message": "Status operasional cafe diperbarui!"})
        except Exception as e:
            db.session.rollback()
            return jsonify({"success": False, "message": str(e)})
    return jsonify({"success": False, "message": "Data cafe tidak ditemukan."})

@owner_bp.route('/api/toggle-jam-operasional', methods=['POST'])
@login_required
@role_required('owner')
def toggle_jam_operasional():
    data = request.json
    jadwal = OperationalHour.query.filter_by(day_of_week=data.get("hari")).first()
    if jadwal:
        jadwal.is_open = data.get("buka")
        try:
            db.session.commit()
            return jsonify({"success": True, "message": "Status jam operasional diperbarui!"})
        except Exception as e:
            db.session.rollback()
            return jsonify({"success": False, "message": str(e)})
    return jsonify({"success": False, "message": "Data hari tidak ditemukan"})